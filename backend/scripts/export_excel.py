import sqlite3
import argparse
from pathlib import Path
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import sys

BASE_DIR = Path(__file__).resolve().parent.parent
ROOT_DIR = BASE_DIR.parent
if str(BASE_DIR / "src") not in sys.path:
    sys.path.insert(0, str(BASE_DIR / "src"))

from market_bot.core.bank_normalizer import normalize_banks

DEFAULT_DB = BASE_DIR / "data" / "market_reports.db"
def style_header_row(ws, row_idx=1, fill_color="006B3F", font_color="FFFFFF"):
    fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    font = Font(name="Calibri", size=11, bold=True, color=font_color)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="medium", color="004D2E"),
        bottom=Side(style="medium", color="004D2E"),
    )

    for cell in ws[row_idx]:
        cell.fill = fill
        cell.font = font
        cell.alignment = align
        cell.border = border
    ws.row_dimensions[row_idx].height = 28


def auto_fit_columns(ws, max_width=50):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if cell.hyperlink:
                val_str = "Google Maps Link"
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), max_width)


def export_to_excel(db_path: Path, output_file: Path):
    if not db_path.exists():
        print(f"Error: Database not found: {db_path}")
        return

    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    cur = conn.cursor()

    wb = openpyxl.Workbook()

    # =========================================================================
    # TAB 1: All Merchant Visits (Primary Dataset)
    # =========================================================================
    ws1 = wb.active
    ws1.title = "All_Merchant_Visits"
    ws1.views.sheetView[0].showGridLines = True

    headers1 = [
        "Visit ID",
        "Date",
        "Time (UTC)",
        "Sales Rep",
        "Merchant Name",
        "Telephone",
        "Business Category",
        "District (Khan)",
        "Status Code",
        "KHQR Stage",
        "Normalized Banks",
        "Deployed Channels",
        "Has CMCB",
        "Has ABA",
        "Has ACLEDA",
        "Has Wing",
        "Has Canadia",
        "Has KB Prasac",
        "Has Sathapana",
        "Has Other Bank",
        "Priority",
        "Target Product",
        "Latitude",
        "Longitude",
        "Google Maps Link",
        "Photos Count",
        "Remark / Notes",
        "Raw Bank Input",
        "Raw Tel",
        "Raw Status",
    ]
    ws1.append(headers1)
    style_header_row(ws1, 1, fill_color="006B3F")

    cur.execute("SELECT * FROM merchant_reports ORDER BY id ASC;")
    reports = cur.fetchall()

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for idx, r in enumerate(reports, start=2):
        report_time_str = r["report_time"] or ""
        date_part = report_time_str[:10]
        time_part = report_time_str[11:19]

        map_link = r["map_url"]
        bank_norm = normalize_banks(r["bank"])
        row_values = [
            r["id"],
            date_part,
            time_part,
            r["sender_name"],
            r["merchant_name"],
            r["tel_clean"],
            r["business_clean"],
            r["district_khan"],
            r["status_clean"],
            r["khqr_stage"],
            bank_norm.canonical_summary,
            bank_norm.products_summary,
            bank_norm.has_cmcb,
            bank_norm.has_aba,
            bank_norm.has_acleda,
            bank_norm.has_wing,
            bank_norm.has_canadia,
            bank_norm.has_kb_prasac,
            bank_norm.has_sathapana,
            bank_norm.has_other_bank,
            r["potential_level"],
            r["potential_product"],
            r["latitude"],
            r["longitude"],
            "Open Google Maps",
            r["photo_count"],
            r["remark"],
            r["bank"],
            r["tel"],
            r["status"],
        ]
        ws1.append(row_values)

        # Style data cells
        for col_idx in range(1, len(headers1) + 1):
            cell = ws1.cell(row=idx, column=col_idx)
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin_border
            if idx % 2 == 0:
                cell.fill = zebra_fill

            # Alignments
            # Center-align IDs, dates, codes, and boolean flags
            if col_idx in [1, 2, 3, 9, 10, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21, 22, 26]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [23, 24]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            # Hyperlink for Google Maps (Column 25)
            if col_idx == 25 and map_link:
                cell.hyperlink = map_link
                cell.font = Font(name="Calibri", size=10, color="005E9E", underline="single")
                cell.alignment = Alignment(horizontal="center", vertical="center")

    ws1.freeze_panes = "A2"
    auto_fit_columns(ws1)

    # =========================================================================
    # TAB 2: Summary By District (Khan Breakdown)
    # =========================================================================
    ws2 = wb.create_sheet(title="Summary_By_District")
    ws2.views.sheetView[0].showGridLines = True

    headers2 = [
        "District (Khan)",
        "Total Visits",
        "CMCB Deployed",
        "In Pipeline",
        "Competitor Targets",
        "Coverage Status",
    ]
    ws2.append(headers2)
    style_header_row(ws2, 1, fill_color="005E9E")

    cur.execute(
        """
        SELECT 
            district_khan,
            count(*) as total_visits,
            sum(CASE WHEN khqr_stage IN ('CONVERTED', 'EXISTING') THEN 1 ELSE 0 END) as deployed_cnt,
            sum(CASE WHEN khqr_stage = 'PIPELINE' THEN 1 ELSE 0 END) as pipeline_cnt,
            sum(CASE WHEN khqr_stage = 'TARGET' THEN 1 ELSE 0 END) as target_cnt
        FROM merchant_reports
        GROUP BY district_khan
        ORDER BY total_visits DESC;
        """
    )
    for idx, r in enumerate(cur.fetchall(), start=2):
        visits = r["total_visits"]
        coverage = "🔥 Hotspot (20+ visits)" if visits >= 20 else "🟢 Active Corridor (5-19)" if visits >= 5 else "🟡 Moderate (1-4)"
        ws2.append([
            r["district_khan"],
            visits,
            r["deployed_cnt"],
            r["pipeline_cnt"],
            r["target_cnt"],
            coverage,
        ])
        for c in range(1, len(headers2) + 1):
            cell = ws2.cell(row=idx, column=c)
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin_border
            if c > 1 and c < 6:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    ws2.freeze_panes = "A2"
    auto_fit_columns(ws2)

    # =========================================================================
    # TAB 3: Summary By Sales Rep
    # =========================================================================
    ws3 = wb.create_sheet(title="Summary_By_Sales_Rep")
    ws3.views.sheetView[0].showGridLines = True

    headers3 = [
        "Sales Representative",
        "Total Visits",
        "CMCB Deployed",
        "In Pipeline",
        "Competitor Targets",
        "Top Active District",
    ]
    ws3.append(headers3)
    style_header_row(ws3, 1, fill_color="102A5E")

    cur.execute(
        """
        SELECT 
            sender_name,
            count(*) as total_visits,
            sum(CASE WHEN khqr_stage IN ('CONVERTED', 'EXISTING') THEN 1 ELSE 0 END) as deployed_cnt,
            sum(CASE WHEN khqr_stage = 'PIPELINE' THEN 1 ELSE 0 END) as pipeline_cnt,
            sum(CASE WHEN khqr_stage = 'TARGET' THEN 1 ELSE 0 END) as target_cnt
        FROM merchant_reports
        GROUP BY sender_name
        ORDER BY total_visits DESC;
        """
    )
    reps = cur.fetchall()
    for idx, r in enumerate(reps, start=2):
        sender = r["sender_name"]
        cur.execute(
            """
            SELECT district_khan, count(*) as c 
            FROM merchant_reports 
            WHERE sender_name = ? 
            GROUP BY district_khan 
            ORDER BY c DESC LIMIT 1;
            """,
            (sender,),
        )
        top_khan_row = cur.fetchone()
        top_khan = f"Khan {top_khan_row['district_khan']} ({top_khan_row['c']})" if top_khan_row else "N/A"

        ws3.append([
            sender,
            r["total_visits"],
            r["deployed_cnt"],
            r["pipeline_cnt"],
            r["target_cnt"],
            top_khan,
        ])
        for c in range(1, len(headers3) + 1):
            cell = ws3.cell(row=idx, column=c)
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin_border
            if c in [2, 3, 4, 5]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    ws3.freeze_panes = "A2"
    auto_fit_columns(ws3)

    # =========================================================================
    # TAB 4: Competitor Bank Share
    # =========================================================================
    ws4 = wb.create_sheet(title="Competitor_Bank_Share")
    ws4.views.sheetView[0].showGridLines = True

    headers4 = [
        "Banking Institution",
        "Merchants with Bank",
        "Market Penetration %",
    ]
    ws4.append(headers4)
    style_header_row(ws4, 1, fill_color="006B3F")

    cur.execute("SELECT count(*) FROM merchant_reports;")
    total_m = cur.fetchone()[0] or 1

    cur.execute("SELECT bank FROM merchant_reports;")
    all_banks_rows = cur.fetchall()

    # Dynamically tally all normalized Cambodian banking institutions across the dataset
    bank_counts = {}
    channel_counts = {"KHQR": 0, "POS": 0, "LOAN": 0}

    for row in all_banks_rows:
        norm = normalize_banks(row["bank"])
        for inst in norm.institutions:
            bank_counts[inst] = bank_counts.get(inst, 0) + 1
        for prod in norm.products:
            if prod in channel_counts:
                channel_counts[prod] += 1

    # Sort by market share descending
    sorted_banks = sorted(bank_counts.items(), key=lambda x: -x[1])

    for idx, (b_name, b_count) in enumerate(sorted_banks, start=2):
        pct = (b_count / total_m) * 100
        ws4.append([b_name, b_count, f"{pct:.1f}%"])
        for c in range(1, 4):
            cell = ws4.cell(row=idx, column=c)
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin_border
            if c in [2, 3]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
    auto_fit_columns(ws4)

    # =========================================================================
    # TAB 5: Compliance Errors (DLQ)
    # =========================================================================
    ws5 = wb.create_sheet(title="Compliance_Errors_DLQ")
    ws5.views.sheetView[0].showGridLines = True

    headers5 = [
        "DLQ ID",
        "Telegram Msg ID",
        "Timestamp",
        "Sales Rep",
        "Error Category",
        "Reason / Issue",
        "Raw Content Preview",
    ]
    ws5.append(headers5)
    style_header_row(ws5, 1, fill_color="D71920")

    cur.execute("SELECT * FROM dead_letter_queue ORDER BY id ASC;")
    dlq_rows = cur.fetchall()

    for idx, d in enumerate(dlq_rows, start=2):
        ws5.append([
            d["id"],
            d["msg_id"],
            d["msg_date"],
            d["sender_name"],
            d["message_type"],
            d["reason"],
            (d["raw_content"] or "")[:150],
        ])
        for c in range(1, len(headers5) + 1):
            cell = ws5.cell(row=idx, column=c)
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin_border
            if c in [1, 2, 3, 5]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    ws5.freeze_panes = "A2"
    auto_fit_columns(ws5)

    # Save Workbook
    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)
    print(f"\n[✓] Excel file generated successfully: {output_file.resolve()}")
    print(f"    - Total Merchant Visits Exported: {len(reports)}")
    print(f"    - Total DLQ Items Exported      : {len(dlq_rows)}")
    conn.close()


def main():
    parser = argparse.ArgumentParser(description="Export market reports database to Excel workbook.")
    parser.add_argument("--db", type=str, default=None, help="Custom database path")
    parser.add_argument("--out", type=str, default="CMCB_Market_Visit_Reports.xlsx", help="Output Excel file path")
    args = parser.parse_args()

    db_path = Path(args.db) if args.db else DEFAULT_DB
    output_file = Path(args.out)

    export_to_excel(db_path, output_file)


if __name__ == "__main__":
    main()
