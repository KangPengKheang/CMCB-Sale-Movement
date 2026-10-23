import argparse
import json
import re
import sqlite3
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl

BASE_DIR = Path(__file__).resolve().parent.parent
ROOT_DIR = BASE_DIR.parent
sys.path.insert(0, str(BASE_DIR))

from src.market_bot.core.merchant_resolver import resolve_or_create_merchant
from src.market_bot.core.parser import (
    classify_khqr_stage,
    match_khan_boundary,
    normalize_business,
    normalize_phone,
    normalize_status,
    parse_competitor_banks,
)
from src.market_bot.core.rm_resolver import clean_rm_name, resolve_or_create_rm
from src.market_bot.db.models import MerchantReportRecord, SenderInfo
from src.market_bot.db.storage import (
    get_all_merchants,
    get_all_rms,
    get_connection,
    init_db,
    recalculate_rm_aggregates,
    save_dlq_item,
    upsert_merchant_report,
)

ICT = timezone(timedelta(hours=7))
KHANS_GEOJSON_PATH = BASE_DIR / "data" / "boundaries" / "phnom_penh_khans.json"
DEFAULT_TARGET_DB = BASE_DIR / "data" / "market_reports_v2.db"
DEFAULT_SOURCE_EXCEL = ROOT_DIR / "Book1.xlsx"
DEFAULT_LIVE_DB = BASE_DIR / "data" / "market_reports.db"


def extract_merchant_name(raw_text: str, name_col: Any) -> str:
    """
    Extracts the merchant name from Excel column or raw report text.
    Handles standard 'Name: <val>' and un-coloned 'Name <val>' formats line-by-line.
    Returns empty string for blank template forms.
    """
    if name_col is not None:
        s = str(name_col).strip()
        if s and s.upper() not in ["NONE", "N/A", ""]:
            return s

    if not raw_text:
        return ""

    for line in raw_text.splitlines():
        line_clean = line.strip()
        # Look for Name: <val> or Name <val>
        m = re.match(r"^Name\s*[:\s]?\s*(.*)$", line_clean, re.IGNORECASE)
        if m:
            val = m.group(1).strip()
            if val and val.upper() not in ["NONE", "N/A", ""]:
                return val
    return ""


def parse_excel_timestamp(msg_date: Any) -> str:
    """
    Converts Excel message date (assumed ICT / UTC+7 if naive) into UTC ISO format:
    YYYY-MM-DDTHH:MM:SSZ
    """
    if isinstance(msg_date, datetime):
        if msg_date.tzinfo is None:
            dt_utc = msg_date.replace(tzinfo=ICT).astimezone(timezone.utc)
        else:
            dt_utc = msg_date.astimezone(timezone.utc)
        return dt_utc.strftime("%Y-%m-%dT%H:%M:%SZ")
    elif isinstance(msg_date, str) and msg_date.strip():
        s = msg_date.strip()
        # Try parsing ISO strings
        try:
            dt = datetime.fromisoformat(s.replace("Z", "+00:00"))
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=ICT)
            return dt.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
        except Exception:
            return s
    return "2026-09-01T00:00:00Z"


def load_pp_khan_features() -> List[Dict[str, Any]]:
    if not KHANS_GEOJSON_PATH.exists():
        return []
    with open(KHANS_GEOJSON_PATH, "r", encoding="utf-8") as f:
        khans_raw = json.load(f)
    return [
        f for f in khans_raw.get("features", [])
        if f.get("properties", {}).get("hrparent") == "KH12"
    ]


def ingest_dataset(
    target_db_path: Path,
    source_excel_path: Path,
    live_db_path: Optional[Path] = None,
) -> Dict[str, Any]:
    print("=" * 80)
    print(" CMCB OFFICIAL EXCEL DATASET INGESTION & ENTITY RESOLUTION ENGINE")
    print("=" * 80)
    print(f"Target Database : {target_db_path}")
    print(f"Source Excel    : {source_excel_path}")
    print(f"Live Database   : {live_db_path or 'None'}")
    print("-" * 80)

    if not source_excel_path.exists():
        raise FileNotFoundError(f"Source Excel file does not exist: {source_excel_path}")

    # Initialize fresh target database
    if target_db_path.exists():
        target_db_path.unlink()
    init_db(target_db_path)
    conn = get_connection(target_db_path)

    pp_features = load_pp_khan_features()
    print(f"[+] Loaded {len(pp_features)} Phnom Penh Khan polygon boundaries.")

    # 1. Load Excel rows
    print(f"[*] Loading master dataset from {source_excel_path.name}...")
    wb = openpyxl.load_workbook(source_excel_path, data_only=True)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]

    excel_signatures = set()
    all_visits: List[MerchantReportRecord] = []
    dlq_count = 0

    for r in range(2, ws.max_row + 1):
        row = {headers[c - 1]: ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)}
        raw_text = str(row.get("Raw_Text") or "")

        sender_id = int(row.get("Sender_ID") or 0)
        sender_name = str(row.get("Sender_Name") or "").strip()
        iso_time = parse_excel_timestamp(row.get("Message_Date"))

        merchant_name = extract_merchant_name(raw_text, row.get("Name"))
        if not merchant_name:
            dlq_count += 1
            save_dlq_item(
                conn=conn,
                msg_id=r,
                msg_date=iso_time,
                sender_id=sender_id,
                sender_name=sender_name,
                message_type="excel_row",
                reason="Missing or empty merchant name template (blank form)",
                raw_content=raw_text or json.dumps({k: str(v) for k, v in row.items() if v is not None}),
            )
            continue

        lat = float(row.get("Latitude") or 0.0)
        lng = float(row.get("Longitude") or 0.0)
        khan = match_khan_boundary(lng, lat, pp_features)

        tel = str(row.get("Tel") or "")
        tel_clean = normalize_phone(tel)

        business = str(row.get("Business") or "")
        biz_clean = normalize_business(business)

        purpose = str(row.get("Purpose") or "CMCB Products")
        bank = str(row.get("Bank") or "")
        has_aba, has_acleda, has_cmcb, has_other = parse_competitor_banks(bank)

        amount = str(row.get("Amount") or "N/A")
        interest = str(row.get("Interest") or "N/A")
        loan_type = str(row.get("Loan_Type") or "N/A")
        tenure = str(row.get("Tenure") or "N/A")
        maturity = str(row.get("Maturity") or "N/A")

        status = str(row.get("Status") or "OP")
        status_clean = normalize_status(status)

        potential_level = str(row.get("Potential_Level") or "L")
        potential_product = str(row.get("Potential_Product") or "KHQR")
        remark = str(row.get("Remark") or "")

        khqr_stage = classify_khqr_stage(status_clean, remark, bank)
        map_url = f"https://www.google.com/maps?q={lat},{lng}" if lat and lng else ""

        # Use synthetic form_msg_id offset by 1,000,000 to avoid collision with Telegram message IDs
        form_msg_id = 1000000 + r
        location_msg_id = 2000000 + r

        record = MerchantReportRecord(
            report_time=iso_time,
            sender_id=sender_id,
            sender_name=sender_name,
            merchant_name=merchant_name,
            tel=tel,
            tel_clean=tel_clean,
            business=business,
            business_clean=biz_clean,
            purpose=purpose,
            bank=bank,
            has_aba=has_aba,
            has_acleda=has_acleda,
            has_cmcb=has_cmcb,
            has_other_bank=has_other,
            amount=amount,
            interest=interest,
            loan_type=loan_type,
            tenure=tenure,
            maturity=maturity,
            status=status,
            status_clean=status_clean,
            khqr_stage=khqr_stage,
            potential_level=potential_level,
            potential_product=potential_product,
            remark=remark,
            latitude=lat,
            longitude=lng,
            district_khan=khan,
            map_url=map_url,
            location_time=iso_time,
            location_diff_seconds=0,
            form_msg_id=form_msg_id,
            location_msg_id=location_msg_id,
            photo_count=1 if row.get("Has_Image") else 0,
            photo_msg_ids="",
            raw_text=raw_text,
        )

        sig = f"{iso_time[:16]}_{sender_id}"
        excel_signatures.add(sig)
        all_visits.append(record)

    print(f"[+] Loaded {len(all_visits)} valid visits from Excel ({dlq_count} routed to DLQ).")

    # 2. Incremental live SQLite extraction strictly after HAILA Hot Pot & BBQ
    # HAILA Hot Pot & BBQ by Mao Theary (8989628727) is the latest record in Book1.xlsx (2026-09-10T05:56:15Z UTC / 12:56 ICT)
    HAILA_CUTOFF_UTC = "2026-09-10T05:56:15Z"
    incremental_count = 0
    if live_db_path and live_db_path.exists():
        print(f"[*] Extracting incremental live visits strictly after HAILA ({HAILA_CUTOFF_UTC}) from {live_db_path.name}...")
        live_conn = sqlite3.connect(live_db_path)
        live_conn.row_factory = sqlite3.Row
        live_cur = live_conn.cursor()
        live_rows = live_cur.execute(
            "SELECT * FROM merchant_reports WHERE report_time > ? ORDER BY report_time ASC;",
            (HAILA_CUTOFF_UTC,),
        ).fetchall()

        for lr in live_rows:
            incremental_count += 1
            rec = MerchantReportRecord(
                report_time=lr["report_time"],
                sender_id=lr["sender_id"],
                sender_name=lr["sender_name"],
                merchant_name=lr["merchant_name"],
                tel=lr["tel"] or "",
                tel_clean=lr["tel_clean"] or "",
                business=lr["business"] or "",
                business_clean=lr["business_clean"] or "",
                purpose=lr["purpose"] or "",
                bank=lr["bank"] or "",
                has_aba=lr["has_aba"] or 0,
                has_acleda=lr["has_acleda"] or 0,
                has_cmcb=lr["has_cmcb"] or 0,
                has_other_bank=lr["has_other_bank"] or 0,
                amount=lr["amount"] or "N/A",
                interest=lr["interest"] or "N/A",
                loan_type=lr["loan_type"] or "N/A",
                tenure=lr["tenure"] or "N/A",
                maturity=lr["maturity"] or "N/A",
                status=lr["status"] or "OP",
                status_clean=lr["status_clean"] or "OP",
                khqr_stage=lr["khqr_stage"] or "TARGET",
                potential_level=lr["potential_level"] or "L",
                potential_product=lr["potential_product"] or "KHQR",
                remark=lr["remark"] or "",
                latitude=lr["latitude"] or 0.0,
                longitude=lr["longitude"] or 0.0,
                district_khan=lr["district_khan"] or "Other / Outside PP",
                map_url=lr["map_url"] or "",
                location_time=lr["location_time"] or lr["report_time"],
                location_diff_seconds=lr["location_diff_seconds"] or 0,
                form_msg_id=lr["form_msg_id"],
                location_msg_id=lr["location_msg_id"],
                photo_count=lr["photo_count"] or 0,
                photo_msg_ids=lr["photo_msg_ids"] or "",
                raw_text=lr["raw_text"] or "",
            )
            all_visits.append(rec)
        live_conn.close()
        print(f"[+] Loaded {incremental_count} incremental live visits strictly after HAILA from {live_db_path.name}.")

    # 3. Sort chronologically
    print(f"[*] Sorting all {len(all_visits)} combined records in strict chronological order...")
    all_visits.sort(key=lambda v: (v.report_time, v.sender_id))

    # 4. Ingestion & Entity Resolution
    print("[*] Processing visits through RM & Merchant Entity Resolution...")
    for idx, v in enumerate(all_visits, start=1):
        # Extract username if sender_name has @username
        username = None
        u_match = re.search(r"@([A-Za-z0-9_]+)", v.sender_name)
        if u_match:
            username = u_match.group(1)

        sender_info = SenderInfo(
            sender_id=v.sender_id,
            sender_name=clean_rm_name(v.sender_name),
            username=username,
        )
        rm_id = resolve_or_create_rm(conn, sender_info, v.report_time)
        v.rm_id = rm_id

        m_id = resolve_or_create_merchant(conn, v)
        v.merchant_id = m_id

        upsert_merchant_report(conn, v)

    # 5. Recalculate RM Aggregates for all RMs
    print("[*] Recalculating RM lifetime metrics and habits...")
    cur = conn.cursor()
    cur.execute("SELECT id FROM rms;")
    rm_ids = [r[0] for r in cur.fetchall()]
    for r_id in rm_ids:
        recalculate_rm_aggregates(conn, r_id)

    # 6. Verification and Summary
    cur.execute("SELECT count(*) FROM merchant_reports;")
    total_reports = cur.fetchone()[0]

    cur.execute("SELECT count(*) FROM merchants;")
    total_merchants = cur.fetchone()[0]

    cur.execute("SELECT count(*) FROM rms;")
    total_rms = cur.fetchone()[0]

    cur.execute("SELECT count(*) FROM dead_letter_queue;")
    total_dlq = cur.fetchone()[0]

    cur.execute("SELECT count(*) FROM merchants WHERE total_visits > 1;")
    multi_visit_merchants = cur.fetchone()[0]

    cur.execute("SELECT count(*) FROM merchants WHERE is_onboarded = 1;")
    onboarded_merchants = cur.fetchone()[0]

    print("\n" + "=" * 80)
    print(" INGESTION & STAR SCHEMA RESOLUTION SUMMARY")
    print("=" * 80)
    print(f"Total Merchant Reports Ingested  : {total_reports}")
    print(f"Unique Merchant Entities Resolved: {total_merchants}")
    print(f"Active Relationship Managers (RM): {total_rms}")
    print(f"Dead-Letter Queue Anomalies      : {total_dlq}")
    print(f"Multi-Visit / Re-visited Stores  : {multi_visit_merchants}")
    print(f"CMCB Onboarded / Active Merchants: {onboarded_merchants}")
    print("=" * 80)

    # RM Performance Leaderboard
    rms = get_all_rms(conn, sort_by="visits")
    print("\n--- RM Performance Leaderboard ---")
    print(f"{'#':<3} | {'RM Name':<30} | {'Visits':<6} | {'Stores':<6} | {'Onboarded':<9} | {'Conv. Rate':<10} | {'Velocity'}")
    print("-" * 90)
    for idx, rm in enumerate(rms, start=1):
        conv_str = f"{rm['conversion_rate']:.1f}%"
        velo_str = f"{rm['daily_visit_velocity']:.1f} / day"
        print(f"{idx:02d}  | {rm['canonical_name'][:28]:<30} | {rm['total_visits_logged']:<6} | {rm['unique_merchants_reached']:<6} | {rm['onboarded_count']:<9} | {conv_str:<10} | {velo_str}")

    conn.close()

    return {
        "total_reports": total_reports,
        "total_merchants": total_merchants,
        "total_rms": total_rms,
        "total_dlq": total_dlq,
        "multi_visit_merchants": multi_visit_merchants,
        "onboarded_merchants": onboarded_merchants,
    }


def main():
    parser = argparse.ArgumentParser(description="Ingest master Book1.xlsx dataset and incremental SQLite reports.")
    parser.add_argument("--target-db", type=str, default=str(DEFAULT_TARGET_DB), help="Path to target output SQLite database")
    parser.add_argument("--source-excel", type=str, default=str(DEFAULT_SOURCE_EXCEL), help="Path to master Book1.xlsx file")
    parser.add_argument("--live-db", type=str, default=str(DEFAULT_LIVE_DB), help="Path to existing SQLite DB for incremental visits")
    args = parser.parse_args()

    target_db = Path(args.target_db)
    source_excel = Path(args.source_excel)
    live_db = Path(args.live_db) if args.live_db else None

    ingest_dataset(target_db_path=target_db, source_excel_path=source_excel, live_db_path=live_db)


if __name__ == "__main__":
    main()
